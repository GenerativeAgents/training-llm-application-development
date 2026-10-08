"""xlsx の図形(drawing XML)をシートごとに、pptx をスライドごとに PNG に描く。マルチモーダル LLM に読ませるための画像化。

LibreOffice で画像化すると、片端だけ図形に接続された折れ線コネクタの経路を誤ることがあったので、
自前で描く。Excel / PowerPoint と同じ見た目は目指さない。箱・ラベル・矢印の向き・チェックボックスの
状態が読めればよい。

図形の中身(prstGeom / solidFill / ln / コネクタ / 文字)は xlsx も pptx も同じ DrawingML(a:)なので、
`Collector` が共通に集める。違うのは外側の名前空間(xdr: / p:)と、図形の位置の出し方だけ:

- xlsx: 位置はセルのアンカー(from/to)から出す。xfrm の off は古いことがあるので使わない。
  フォームコントロールのチェックボックスは xl/ctrlProps/ の checked を読む。
  セルは文字列と塗りだけ下地として描く(罫線は描かない)。帳票型のシートではチェックボックスの
  意味が行ラベルで決まるため。表の内容そのものは structured 抽出器がテキストで持っている
- pptx: 位置は xfrm(EMU の絶対座標)。プレースホルダで xfrm が無いものはレイアウト → マスターから継ぐ。
  表(graphicFrame の a:tbl)はセルを矩形 + 文字として描く
- グラフ(graphicFrame の c:chart。xlsx も pptx も同じ): 描かずに、枠と「グラフ: 題」だけを描く。
  値はグラフの XML に入っているので chart_lines でテキストにする(抽出器が本文に入れる)
- 貼られた画像のうち EMF / WMF(Excel / PowerPoint のグラフを貼るとこの形式になる)は Pillow が読めないので、
  LibreOffice に PNG を描かせてから貼る(metafile_png)
- 共通: グループ内の図形はグループの chOff/chExt から親の枠へ比例変換する。
  コネクタは stCxn/endCxn(接続先の図形と接続点番号)から経路を組み立てる。保存された枠は
  接続済みの端では信用しない(Excel は表示時に計算し直しているため)。
  文字は段落ごとに改行し、pPr の algn / bodyPr の anchor / normAutofit の fontScale を反映する
  (ポンチ絵は文字量が多く、全部中央寄せにすると箱からはみ出して隣と重なる)

    uv run python -m app.advanced_rag.extractors.render_drawing <xlsx|pptx> --out DIR              # 図形のあるシート / スライドを全部
    uv run python -m app.advanced_rag.extractors.render_drawing <xlsx> --sheet NAME --out DIR
    uv run python -m app.advanced_rag.extractors.render_drawing <pptx> --sheet slide08 --out DIR
"""

import argparse
import hashlib
import io
import math
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Iterable, Iterator, Sequence
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, cast

import openpyxl
from openpyxl.cell.cell import Cell as ExcelCell
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageChops, ImageDraw, ImageFont

Point = tuple[float, float]
Box = Sequence[float]  # (x, y, 幅, 高さ)。リストとタプルを受け取る。

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"
CHART_TYPES = {"barChart": "棒グラフ", "bar3DChart": "棒グラフ", "lineChart": "折れ線グラフ", "line3DChart": "折れ線グラフ",
               "pieChart": "円グラフ", "pie3DChart": "円グラフ", "ofPieChart": "円グラフ", "doughnutChart": "ドーナツグラフ",
               "areaChart": "面グラフ", "area3DChart": "面グラフ", "scatterChart": "散布図", "radarChart": "レーダーチャート",
               "bubbleChart": "バブルチャート", "stockChart": "株価チャート", "surfaceChart": "等高線グラフ", "surface3DChart": "等高線グラフ"}
EMU_PX = 9525  # 96dpi
SCALE = 1.5  # 描画倍率(文字を読みやすく)
TITLE_PX = 24  # 画像の上端に題(ファイル名 / シート名)を書く帯の高さ(px, 96dpi)。図形はこの分下にずらして描く
MIN_SHAPES = 3  # これ未満のシート / スライドは描かない(表紙のロゴ枠、題だけのスライドなど)
ELBOW_MARGIN = 20  # U 型コネクタが図形から外へ出る距離(px)
DEFAULT_PT = 9

THEME_ALIAS = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
# prstClr(色名での指定)のうちよく使われるもの。無い名前は色なし扱い
PRESET_COLORS = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
                 "yellow": "FFFF00", "gray": "808080", "grey": "808080", "darkGray": "A9A9A9", "lightGray": "D3D3D3"}
# IPA フォントに無い丸数字(➀〜➉、➊〜➓)は ①〜⑩ で描く(ポンチ絵の矢印ラベルに使われる。豆腐になると LLM が番号を読めない)
GLYPH_FALLBACK = {c: chr(0x2460 + i) for base in (0x2780, 0x278A) for i, c in enumerate(range(base, base + 10))}


def find_font() -> str | None:
    try:
        out = subprocess.run(["fc-list", ":lang=ja", "file"], capture_output=True, text=True).stdout
    except FileNotFoundError:
        return None
    paths = [l.strip().rstrip(":") for l in out.splitlines() if l.strip()]
    # プロポーショナルのゴシック(ipagp)を優先。Office の既定(ＭＳ Ｐゴシック / HGP ゴシック)に字幅が近い
    paths.sort(key=lambda p: (0 if "ipagp" in p.lower() else 1 if "gothic" in p.lower() else 2, p))
    return paths[0] if paths else None


FONT_PATH = find_font()


def font(pt: float, scale: float = SCALE) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    px = max(8, int(pt * 96 / 72 * scale))
    return ImageFont.truetype(FONT_PATH, px) if FONT_PATH else ImageFont.load_default()


# ---------------------------------------------------------------- EMF / WMF


METAFILE_CACHE = Path(__file__).resolve().parents[3] / "data" / "cache" / "metafile"
SOFFICE_PROFILE = Path(__file__).resolve().parents[3] / "data" / "cache" / "soffice"  # 利用者の LibreOffice と別のプロファイル
METAFILE_PX = (3000, 4243)  # LibreOffice に描かせる大きさ(A4 の縦横比)。余白は後で落とすので、図はこれより小さくなる


def metafile_suffix(data: bytes) -> str | None:
    """EMF / WMF ならその拡張子。soffice は拡張子で形式を決めるので、中身から判定して付け直す。"""
    if data[:4] == b"\x01\x00\x00\x00" and data[40:44] == b" EMF":
        return ".emf"
    if data[:4] in (b"\xd7\xcd\xc6\x9a", b"\x01\x00\x09\x00"):
        return ".wmf"
    return None


def metafile_png(data: bytes) -> bytes | None:
    """EMF / WMF を PNG にする。Pillow が開けないので LibreOffice(Draw)に描かせる。

    グラフや図を Excel / PowerPoint に貼るとこの形式になることが多く、そのままではマルチモーダル LLM に渡す画像に何も写らない。
    soffice は 1 ページの Draw 文書として開くため、描かれた部分(白でないところ)だけに切り抜く。
    LibreOffice が無い、EMF / WMF でない、変換に失敗した、のいずれでも None を返す(呼び出し側で画像なしとして扱う)。
    結果は data/cache/metafile/ に残す(同じ画像を何度も変換しない。失敗も空ファイルで覚える)。
    """
    suffix = metafile_suffix(data)
    if not suffix:
        return None
    cache = METAFILE_CACHE / f"{hashlib.sha1(data).hexdigest()}.png"
    if cache.exists():
        return cache.read_bytes() or None
    METAFILE_CACHE.mkdir(parents=True, exist_ok=True)
    opts = '{"PixelWidth":{"type":"long","value":%d},"PixelHeight":{"type":"long","value":%d}}' % METAFILE_PX
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"metafile{suffix}"
        src.write_bytes(data)
        try:
            subprocess.run(["soffice", f"-env:UserInstallation=file://{SOFFICE_PROFILE}", "--headless",
                            "--convert-to", f"png:draw_png_Export:{opts}", "--outdir", tmp, str(src)],
                           capture_output=True, timeout=180, check=True)
        except (OSError, subprocess.SubprocessError):
            cache.write_bytes(b"")
            return None
        out = Path(tmp) / "metafile.png"
        png = crop_margins(out) if out.exists() else None
    cache.write_bytes(png or b"")
    return png


def crop_margins(path: Path) -> bytes | None:
    """白い余白を落として、描かれた部分だけの PNG にする。"""
    im = Image.open(path).convert("RGB")
    box = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()
    if not box:
        return None
    buf = io.BytesIO()
    im.crop(box).save(buf, format="PNG")
    return buf.getvalue()


def open_image(data: bytes) -> Image.Image:
    """画像のバイト列を Pillow で開く。EMF / WMF は PNG にしてから開く。

    Pillow は EMF / WMF のヘッダ(大きさ)は読めるが中身は描けず、`Image.open` は通って `load` で失敗する。
    それでは呼び出し側が「開けた」と勘違いするので、形式で判定して先に変換する。変換できなければ、
    そのまま Pillow に渡して例外にする(呼び出し側が画像なしとして扱う)。
    """
    if metafile_suffix(data):
        png = metafile_png(data)
        if png is not None:
            return Image.open(io.BytesIO(png))
    return Image.open(io.BytesIO(data))


# ---------------------------------------------------------------- package parts (xlsx / pptx 共通)


class Package:
    """OOXML の zip。関係ファイル(rels)の解決、テーマ色、メディアの取り出し。"""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.z = zipfile.ZipFile(path)
        self.names = set(self.z.namelist())
        self.theme: dict[str, str] = {}

    def _rels(self, part: str) -> dict[str, str]:
        if part not in self.names:
            return {}
        out = {}
        for m in re.finditer(r"<Relationship [^>]*/>", self.z.read(part).decode()):
            rid = re.search(r'Id="([^"]*)"', m.group(0))
            target = re.search(r'Target="([^"]*)"', m.group(0))
            if rid and target and target.group(1):  # 外部ハイパーリンクなどで Target が空のことがある(参照先が無いので飛ばす)
                out[rid.group(1)] = target.group(1)
        return out

    @staticmethod
    def _resolve(base: str, target: str) -> str:
        """rels の Target(パートからの相対パス、または / で始まる絶対パス)を zip 内のパス名にする。"""
        if target.startswith("/"):
            return target.lstrip("/")
        parts = base.split("/")[:-1]
        for seg in target.split("/"):
            if seg == "..":
                parts.pop()
            elif seg and seg != ".":
                parts.append(seg)
        return "/".join(parts)

    def _theme_from(self, part: str) -> dict[str, str]:
        if part not in self.names:
            return {}
        t = self.z.read(part).decode()
        return dict(re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)>\s*<a:(?:srgbClr val|sysClr val="\w+" lastClr)="(\w+)"', t))

    def media(self, base: str, rels: dict[str, str], rid: str) -> bytes | None:
        """base(rels を持つパート)から見た相対 Target の画像バイト列。"""
        target = rels.get(rid)
        if not target:
            return None
        part = self._resolve(base, target)
        return self.z.read(part) if part in self.names else None

    def color(self, node: ET.Element | None) -> str | None:
        """solidFill の子(srgbClr / schemeClr)を #RRGGBB に。lumMod / lumOff(明るさの補正)はどちらにも掛ける。"""
        if node is None:
            return None
        clr, hexv = node.find("a:srgbClr", NS), None
        if clr is not None:
            hexv = cast(str, clr.get("val"))
        elif node.find("a:prstClr", NS) is not None:
            clr = cast(ET.Element, node.find("a:prstClr", NS))
            hexv = PRESET_COLORS.get(cast(str, clr.get("val")))
        elif node.find("a:sysClr", NS) is not None:
            clr = cast(ET.Element, node.find("a:sysClr", NS))
            hexv = clr.get("lastClr")
        else:
            clr = node.find("a:schemeClr", NS)
            if clr is not None:
                v = cast(str, clr.get("val"))
                hexv = self.theme.get(THEME_ALIAS.get(v, v))
        if not hexv:
            return None
        clr = cast(ET.Element, clr)
        rgb = tuple(int(hexv[i : i + 2], 16) for i in (0, 2, 4))
        # lumMod / lumOff で明るさだけ近似
        mod = clr.find("a:lumMod", NS)
        off = clr.find("a:lumOff", NS)
        if mod is not None or off is not None:
            m = int(cast(str, mod.get("val"))) / 100000 if mod is not None else 1.0
            o = int(cast(str, off.get("val"))) / 100000 if off is not None else 0.0
            rgb = tuple(min(255, int(c * m + 255 * o)) for c in rgb)
        return "#%02x%02x%02x" % rgb


class Workbook(Package):
    def __init__(self, path: Path) -> None:
        super().__init__(path)
        wb = self.z.read("xl/workbook.xml").decode()
        rels = self._rels("xl/_rels/workbook.xml.rels")
        self.sheets = {}  # name -> worksheet part path
        for m in re.finditer(r"<sheet [^>]*/>", wb):
            name = cast(re.Match[str], re.search(r'name="([^"]+)"', m.group(0))).group(1).replace("&amp;", "&")
            rid = cast(re.Match[str], re.search(r'r:id="([^"]+)"', m.group(0))).group(1)
            self.sheets[name] = "xl/" + rels[rid].lstrip("/").removeprefix("xl/")
        self.theme = self._theme_from("xl/theme/theme1.xml")
        self.mdw = self._max_digit_width()
        self.wb = openpyxl.load_workbook(path, data_only=True)  # 列幅・行高とセルの値(数式は計算済みの値)

    def _max_digit_width(self) -> int:
        """列幅の単位になる標準フォントの数字幅(px)。9pt なら 6、11pt なら 7。"""
        if "xl/styles.xml" not in self.names:
            return 7
        st = self.z.read("xl/styles.xml").decode()
        m = re.search(r"<fonts[^>]*>\s*<font[^>]*>(.*?)</font>", st, re.S)
        sz = re.search(r'<sz val="([\d.]+)"', m.group(1)) if m else None
        pt = float(sz.group(1)) if sz else 11
        return max(5, round(pt * 96 / 72 * 0.5))

    def drawing_for(self, sheet: str) -> tuple[ET.Element, str, dict[str, str], str] | None:
        """(drawing の ElementTree root, drawing のパート名, drawing の rels, worksheet XML)。drawing が無ければ None。"""
        part = self.sheets[sheet]
        rels = self._rels(f"xl/worksheets/_rels/{Path(part).name}.rels")
        ws_xml = self.z.read(part).decode()
        drawing = None
        for rid, target in rels.items():
            if "drawings/drawing" in target and target.endswith(".xml"):
                drawing = self._resolve(part, target)
        if drawing is None or drawing not in self.names:
            return None
        drels = self._rels(f"xl/drawings/_rels/{Path(drawing).name}.rels")
        return ET.fromstring(self.z.read(drawing)), drawing, drels, ws_xml

    def checkbox_states(self, ws_xml: str, part: str) -> dict[str, str]:
        """コントロール名 -> 'Checked' / 'Mixed' / ''(未チェック)"""
        rels = self._rels(f"xl/worksheets/_rels/{Path(part).name}.rels")
        out = {}
        for m in re.finditer(r'<control [^>]*>', ws_xml):
            t = m.group(0)
            name = re.search(r'name="([^"]*)"', t)
            rid = re.search(r'r:id="([^"]+)"', t)
            if not (name and rid and rid.group(1) in rels):
                continue
            cp = self._resolve(part, rels[rid.group(1)])
            if cp not in self.names:
                continue
            x = self.z.read(cp).decode()
            if "<formControlPr" not in x or 'objectType="CheckBox"' not in x:
                continue
            c = re.search(r'checked="(\w+)"', x)
            out[name.group(1)] = c.group(1) if c else ""
        return out


SKIP_PLACEHOLDERS = {"dt", "ftr", "sldNum"}  # 日付・フッタ・スライド番号(図の内容ではない)


class Presentation(Package):
    def __init__(self, path: Path) -> None:
        super().__init__(path)
        pres = self.z.read("ppt/presentation.xml").decode()
        rels = self._rels("ppt/_rels/presentation.xml.rels")
        # スライドの順序は sldIdLst(ファイル名の番号順とは限らない)
        self.slides = []  # [(名前 "slide01", パート名)]
        for i, rid in enumerate(re.findall(r'<p:sldId [^>]*r:id="([^"]+)"', pres), 1):
            self.slides.append((f"slide{i:02d}", self._resolve("ppt/presentation.xml", rels[rid])))
        m = re.search(r'<p:sldSz cx="(\d+)" cy="(\d+)"', pres)
        self.size = (int(m.group(1)) / EMU_PX, int(m.group(2)) / EMU_PX) if m else (960, 540)
        masters = sorted(n for n in self.names if re.match(r"ppt/slideMasters/slideMaster\d+\.xml$", n))
        theme = self._related(masters[0], "theme") if masters else None
        self.theme = self._theme_from(theme or "ppt/theme/theme1.xml")

    def _related(self, part: str, kind: str) -> str | None:
        """part の rels から種類(slideLayout / slideMaster / theme)が kind のパートを 1 つ。"""
        for target in self._rels(f"{Path(part).parent}/_rels/{Path(part).name}.rels").values():
            if f"/{kind}" in target and target.endswith(".xml"):
                return self._resolve(part, target)
        return None

    def placeholder_xfrm(self, part: str, ph: ET.Element) -> dict[str, Any] | None:
        """xfrm を持たないプレースホルダの位置を、レイアウト → マスターの同じプレースホルダから継ぐ。"""
        idx, typ = ph.get("idx"), ph.get("type") or "body"
        layout = self._related(part, "slideLayout")
        chain = [p for p in (layout, self._related(layout, "slideMaster") if layout else None) if p]
        for cand in chain:
            root = ET.fromstring(self.z.read(cand))
            best = None
            for sp in root.iter("{%s}sp" % NS["p"]):
                cph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
                if cph is None:
                    continue
                if idx is not None and cph.get("idx") == idx:
                    best = sp
                    break
                if best is None and (cph.get("type") or "body") == typ:
                    best = sp
            if best is not None:
                f = xfrm(best.find("p:spPr", NS))
                if f:
                    return f
        return None


# ---------------------------------------------------------------- geometry helpers


def cell_grid(ws: Worksheet, mdw: int = 7) -> tuple[list[float], list[float]]:
    """列・行の開始位置(px, 96dpi)。

    openpyxl は <col min max> の範囲を先頭の列にしか持たず、他の列を引くと既定値(13 文字幅)を
    返してしまうので、範囲を自分で展開する。列幅 px は Excel の式
    trunc((256*width + trunc(128/MDW)) / 256 * MDW)。
    """
    dw = ws.sheet_format.defaultColWidth or ws.sheet_format.baseColWidth or 8.38
    dh = ws.sheet_format.defaultRowHeight or 13.5
    widths = {}
    for d in ws.column_dimensions.values():
        if d.width and d.min and d.max:
            for c in range(d.min, min(d.max, 400) + 1):
                widths[c] = d.width
    cols = [0.0]
    for c in range(1, 400):
        w = widths.get(c, dw)
        cols.append(cols[-1] + int((256 * w + int(128 / mdw)) / 256 * mdw))
    rows = [0.0]
    for r in range(1, 400):
        h = ws.row_dimensions[r].height or dh
        rows.append(rows[-1] + h * 96 / 72)
    return cols, rows


def anchor_box(a: ET.Element, cols: list[float], rows: list[float]) -> list[float]:
    def pt(node: ET.Element) -> Point:
        g: Callable[[str], int] = lambda k: int(cast(str, cast(ET.Element, node.find(f"xdr:{k}", NS)).text))
        return cols[min(g("col"), len(cols) - 1)] + g("colOff") / EMU_PX, rows[min(g("row"), len(rows) - 1)] + g("rowOff") / EMU_PX

    x0, y0 = pt(cast(ET.Element, a.find("xdr:from", NS)))
    to = a.find("xdr:to", NS)
    if to is not None:
        x1, y1 = pt(to)
    else:
        ext = cast(ET.Element, a.find("xdr:ext", NS))
        x1, y1 = x0 + int(cast(str, ext.get("cx"))) / EMU_PX, y0 + int(cast(str, ext.get("cy"))) / EMU_PX
    return [x0, y0, x1 - x0, y1 - y0]


def xfrm(el: ET.Element | None) -> dict[str, Any] | None:
    return xfrm_node(el.find("a:xfrm", NS)) if el is not None else None


def xfrm_node(x: ET.Element | None) -> dict[str, Any] | None:
    """xfrm 要素(a:xfrm、または graphicFrame 直下の xdr:xfrm / p:xfrm)を dict に。"""
    if x is None:
        return None
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    if off is None or ext is None:
        return None
    d: dict[str, Any] = dict(x=int(cast(str, off.get("x"))), y=int(cast(str, off.get("y"))), cx=int(cast(str, ext.get("cx"))), cy=int(cast(str, ext.get("cy"))),
             flipH=x.get("flipH") == "1", flipV=x.get("flipV") == "1", rot=int(x.get("rot") or 0) / 60000)
    cho, che = x.find("a:chOff", NS), x.find("a:chExt", NS)
    if cho is not None:
        d["ch"] = (int(cast(str, cho.get("x"))), int(cast(str, cho.get("y"))), int(cast(str, cast(ET.Element, che).get("cx"))), int(cast(str, cast(ET.Element, che).get("cy"))))
    return d


def txt(el: ET.Element) -> str:
    """図形の文字。段落(a:p)ごとに改行する。"""
    paras = ["".join(t.text or "" for t in p.iter("{%s}t" % NS["a"])) for p in el.iter("{%s}p" % NS["a"])]
    return "\n".join(paras).strip("\n")


def font_scale(el: ET.Element) -> float:
    fit = el.find(".//a:bodyPr/a:normAutofit", NS)  # PowerPoint の「はみ出す場合に自動調整」で縮んだ文字
    return int(cast(str, fit.get("fontScale"))) / 100000 if fit is not None and fit.get("fontScale") else 1.0


def text_size_pt(el: ET.Element) -> float:
    szs = [int(cast(str, r.get("sz"))) for r in el.iter("{%s}rPr" % NS["a"]) if r.get("sz")]
    return (szs[0] / 100 if szs else DEFAULT_PT) * font_scale(el)


def paragraphs(el: ET.Element) -> list[tuple[str, float, float | None]]:
    """段落ごとの (文字, pt, 行送り pt または None)。pt は段落内の最初の run の sz(無ければ endParaRPr、それも無ければ直前の段落)。

    ポンチ絵は 1 つの箱の中で見出し行だけ大きい、行送りを固定(lnSpc の spcPts)して詰める、が普通なので段落単位で持つ。
    """
    scale = font_scale(el)
    out: list[tuple[str, float, float | None]] = []
    pt: float = DEFAULT_PT
    for p in el.iter("{%s}p" % NS["a"]):
        text = "".join(t.text or "" for t in p.iter("{%s}t" % NS["a"]))
        szs = [int(cast(str, r.get("sz"))) for tag in ("rPr", "endParaRPr") for r in p.iter("{%s}%s" % (NS["a"], tag)) if r.get("sz")]
        if szs:
            pt = szs[0] / 100
        lh = None
        pts_ = p.find("a:pPr/a:lnSpc/a:spcPts", NS)
        pct = p.find("a:pPr/a:lnSpc/a:spcPct", NS)
        if pts_ is not None:
            lh = int(cast(str, pts_.get("val"))) / 100
        elif pct is not None:
            lh = pt * 1.2 * int(cast(str, pct.get("val"))) / 100000
        out.append((text, pt * scale, lh * scale if lh else None))
    return out


def text_layout(el: ET.Element) -> tuple[str, str, bool, bool, tuple[float, float]]:
    """(横位置, 縦位置, 折り返すか, 縦書きか, (左右の余白 px, 上下の余白 px))。

    指定が無ければ左上・折り返し(PowerPoint のテキストボックスの既定)。余白は指定があるときだけ使う。
    """
    body = el.find(".//a:bodyPr", NS)
    valign = {"ctr": "middle", "b": "bottom"}.get(cast(str, body.get("anchor") if body is not None else None), "top")
    wrap = body is None or body.get("wrap") != "none"
    vert = body is not None and (body.get("vert") or "horz") != "horz"
    ppr = el.find(".//a:p/a:pPr", NS)
    align = {"ctr": "center", "r": "right"}.get(cast(str, ppr.get("algn") if ppr is not None else None), "left")
    ins = (2.0, 1.0)
    if body is not None and (body.get("lIns") or body.get("tIns")):
        ins = (int(body.get("lIns") or 91440) / EMU_PX, int(body.get("tIns") or 45720) / EMU_PX)
    return align, valign, wrap, vert, ins


def text_fields(el: ET.Element) -> dict[str, Any]:
    """図形(または表のセル)の文字に関する項目。items の要素に混ぜる。"""
    align, valign, wrap, vert, ins = text_layout(el)
    return dict(text=txt(el), pt=text_size_pt(el), paras=paragraphs(el), align=align, valign=valign, wrap=wrap, vert=vert, ins=ins)


# ---------------------------------------------------------------- charts


def _chart_text(el: ET.Element | None) -> str:
    """c:tx / c:title の文字(リッチテキストの a:t、またはセル参照のキャッシュ c:v)。"""
    if el is None:
        return ""
    ts = [t.text or "" for t in el.iter("{%s}t" % NS["a"])]
    if ts:
        return "".join(ts).strip()
    return " ".join((v.text or "").strip() for v in el.iter("{%s}v" % NS["c"])).strip()


def format_number(text: str, code: str | None) -> str:
    """数値のキャッシュ(倍精度をそのまま書いた `4.9000000000000002E-2` など)を、表示形式(formatCode)に近い見た目にする。

    よく使う形だけ扱う: パーセント(`0.0%`)、小数の桁数(`0.00`)、桁区切り(`#,##0`)、前後の文字(`0"円"`)、日付(`yyyy/m/d`)。
    それ以外(`General` など)は有効数字 15 桁に丸めて、倍精度の端数だけ落とす。数値でなければそのまま。
    """
    try:
        v = float(text)
    except ValueError:
        return text
    code = (code or "").split(";")[0]  # 正の数の書式だけ見る
    code = re.sub(r"\[[^\]]*\]", "", code)  # [Red] や [$-411] など
    bare = re.sub(r'"[^"]*"', "", code)
    if not re.search(r"[0#?]", bare):
        if re.search(r"[ymd]", bare, re.I):  # 日付: Excel のシリアル値(1900 年起点。1900-02-29 のずれは 1899-12-30 起点で吸収)
            try:
                return (datetime(1899, 12, 30) + timedelta(days=v)).date().isoformat()
            except (OverflowError, ValueError):
                return text
        return f"{v:.15g}"
    if "%" in bare:
        v *= 100
    frac = re.search(r"\.([0#?]+)", bare)
    digits = len(frac.group(1)) if frac else 0
    number = f"{v:,.{digits}f}" if "," in bare.split(".")[0] else f"{v:.{digits}f}"
    first = cast(re.Match[str], re.search(r"[0#?]", code)).start()
    prefix = "".join(re.findall(r'"([^"]*)"', code[:first]))
    suffix = "".join(re.findall(r'"([^"]*)"', code[first:]))
    return prefix + number + ("%" if "%" in bare else "") + suffix


def _chart_points(el: ET.Element | None) -> dict[int, str]:
    """c:cat / c:val / c:xVal / c:yVal のキャッシュ値を 点の番号 → 文字 で。多段の項目(multiLvl)は段を / でつなぐ。

    数値のキャッシュ(c:numCache)は、表示形式(キャッシュの c:formatCode、点ごとの formatCode 属性が優先)で整える。
    """
    if el is None:
        return {}
    levels = el.findall(".//c:lvl", NS)
    groups = levels if levels else [el]
    code_el = el.find(".//c:numCache/c:formatCode", NS)
    code = code_el.text if code_el is not None else None
    numeric = el.find(".//c:numCache", NS) is not None
    out: dict[int, list[str]] = {}
    for g in reversed(groups):  # multiLvl は内側の段から並んでいるので外側を先に
        for pt in g.iter("{%s}pt" % NS["c"]):
            v = pt.find("c:v", NS)
            if v is not None and v.text is not None:
                text = format_number(v.text.strip(), pt.get("formatCode") or code) if numeric else v.text.strip()
                out.setdefault(int(cast(str, pt.get("idx"))), []).append(text)
    return {k: " / ".join(v) for k, v in out.items()}


def chart_data(xml: bytes) -> dict[str, Any] | None:
    """グラフのパート(chartN.xml)から、題・種類・軸の題・系列(名前と (項目, 値) の並び)を取る。値はキャッシュ(c:v)を使う。"""
    root = ET.fromstring(xml)
    chart = root.find("c:chart", NS)
    if chart is None:
        return None
    title = _chart_text(chart.find("c:title", NS))
    plot = chart.find("c:plotArea", NS)
    if plot is None:
        return None
    kinds, series = [], []
    for grp in plot:
        tag = grp.tag.split("}")[1]
        if not tag.endswith("Chart"):
            continue
        kind = CHART_TYPES.get(tag, tag)
        bar_dir = grp.find("c:barDir", NS)
        if tag.startswith("bar") and bar_dir is not None and bar_dir.get("val") == "bar":
            kind = "横棒グラフ"
        kinds.append(kind)
        for ser in grp.findall("c:ser", NS):
            cats = _chart_points(ser.find("c:cat", NS)) or _chart_points(ser.find("c:xVal", NS))
            vals = _chart_points(ser.find("c:val", NS)) or _chart_points(ser.find("c:yVal", NS))
            points = [(cats.get(i, str(i + 1)), vals[i]) for i in sorted(vals)]
            series.append({"name": _chart_text(ser.find("c:tx", NS)), "points": points})
    axes = [t for t in (_chart_text(ax.find("c:title", NS)) for ax in plot if ax.tag.split("}")[1].endswith("Ax")) if t]
    return {"title": title, "kinds": list(dict.fromkeys(kinds)), "axes": axes, "series": series}


def chart_lines(data: dict[str, Any]) -> list[str]:
    """chart_data をテキストの行にする: `グラフ: 題(種類)`、`軸: …`、系列ごとに `系列名: 項目 値 / 項目 値 …`。

    値はグラフの XML のキャッシュ(c:v)から取る。Excel / PowerPoint が保存したファイルには必ずあるが、
    セル範囲の参照だけを書くツールもある(openpyxl など)。そのときは題と軸だけになる。
    """
    kinds = "・".join(data["kinds"]) or "グラフ"
    lines = [f"グラフ: {data['title']}({kinds})" if data["title"] else f"グラフ: ({kinds})"]
    if data["axes"]:
        lines.append("軸: " + " / ".join(data["axes"]))
    for i, ser in enumerate(data["series"], 1):
        if ser["points"]:  # 値のキャッシュが無い系列(Excel 以外のツールで作ったファイル)は出さない
            name = ser["name"] or f"系列{i}"
            lines.append(f"{name}: " + " / ".join(f"{c} {v}" for c, v in ser["points"]))
    return lines


# ---------------------------------------------------------------- collect items


THEME_INDEX = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]


def cell_fill(cell: ExcelCell | MergedCell, theme: dict[str, str]) -> str | None:
    if not cell.fill or cell.fill.fill_type != "solid":
        return None
    fg = cell.fill.fgColor
    rgb = None
    if fg.type == "rgb" and isinstance(fg.rgb, str) and fg.rgb != "00000000":
        rgb = fg.rgb[-6:]
    elif fg.type == "theme" and 0 <= fg.theme < len(THEME_INDEX):
        rgb = theme.get(THEME_INDEX[fg.theme])
    elif fg.type == "indexed":
        from openpyxl.styles.colors import COLOR_INDEX
        if 0 <= fg.indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[fg.indexed][-6:]
    if not rgb:
        return None
    r, g, b = (int(rgb[i : i + 2], 16) for i in (0, 2, 4))
    tint = fg.tint or 0
    if tint > 0:
        r, g, b = (int(c + (255 - c) * tint) for c in (r, g, b))
    elif tint < 0:
        r, g, b = (int(c * (1 + tint)) for c in (r, g, b))
    return "#%02x%02x%02x" % (r, g, b)


def cell_items(ws: Worksheet, cols: list[float], rows: list[float], theme: dict[str, str]) -> list[dict[str, Any]]:
    """値のあるセル(結合セルはその範囲)を、位置・文字列・塗り色で返す。"""
    merged: dict[tuple[int, int], tuple[int, int] | None] = {}
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    merged[(r, c)] = None
    out = []
    for row in ws.iter_rows():
        for cell in row:
            key = (cell.row, cell.column)
            if merged.get(key, key) is None or cell.row >= len(rows) - 1 or cell.column >= len(cols) - 1:
                continue
            r1, c1 = merged.get(key) or key
            r1, c1 = min(r1, len(rows) - 2), min(c1, len(cols) - 2)
            fill = cell_fill(cell, theme)
            if fill in ("#ffffff", "#FFFFFF"):
                fill = None
            v = cell.value
            text = "" if v is None else (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v))
            if not text and not fill:
                continue
            box = [cols[cell.column - 1], rows[cell.row - 1], cols[c1] - cols[cell.column - 1], rows[r1] - rows[cell.row - 1]]
            align = {"center": "center", "right": "right"}.get(cell.alignment.horizontal, "left")
            pt = cell.font.sz or DEFAULT_PT
            wrap = bool(cell.alignment.wrap_text) or merged.get(key) is not None
            out.append(dict(box=box, text=text, fill=fill, align=align, pt=float(pt), wrap=wrap))
    return out


class Collector:
    """drawing XML(xlsx は xdr:、pptx は p:)から図形を位置つきで集める。図形の中身は同じ DrawingML なので共通。

    ns: 外側の名前空間の接頭辞。pkg: 色とメディアの解決。base: メディアの相対パスの基準になるパート名。
    checks: xlsx のチェックボックス状態。inherit: xfrm の無い図形の位置を補う関数(pptx のプレースホルダ用)。
    """

    def __init__(self, ns: str, pkg: Package, base: str, rels: dict[str, str], checks: dict[str, str] | None = None, inherit: Callable[[ET.Element], dict[str, Any] | None] | None=None) -> None:
        self.ns, self.pkg, self.base, self.rels = ns, pkg, base, rels
        self.checks = checks or {}
        self.inherit = inherit
        self.items: list[dict[str, Any]] = []

    def q(self, tag: str) -> str:
        return f"{self.ns}:{tag}"

    def emit(self, el: ET.Element, kind: str, box: Box, f: dict[str, Any] | None) -> None:
        cnv = cast(ET.Element, el.find(f"./*/{self.q('cNvPr')}", NS))
        sppr = el.find(self.q("spPr"), NS)
        it = dict(id=cnv.get("id"), name=cnv.get("name") or "", kind=kind, box=box,
                  flipH=f["flipH"] if f else False, flipV=f["flipV"] if f else False, rot=f["rot"] if f else 0,
                  prst=None, cust=None, fill=None, line=True, lncolor="black", dash=False, head=False, tail=False,
                  adj={}, stCxn=None, endCxn=None, image=None, check=None, **text_fields(el))
        style = el.find(self.q("style"), NS)
        # 文字色: 最初の文字列の書式、無ければ段落末の書式、無ければ図形のスタイル(fontRef)。どれも無ければ黒
        run_fill = el.find(".//a:r/a:rPr/a:solidFill", NS)
        if run_fill is None:
            run_fill = el.find(".//a:endParaRPr/a:solidFill", NS)
        font_ref = style.find("a:fontRef", NS) if style is not None else None
        if run_fill is not None:  # 文字に色の指定があれば、読めない形式でもスタイルの色(白など)には落とさない
            it["color"] = self.pkg.color(run_fill) or "black"
        else:
            it["color"] = (self.pkg.color(font_ref) if font_ref is not None else None) or "black"
        if style is not None:
            # スタイル(テーマの図形の書式)の線・塗りを既定にする。spPr に指定があればそちらが優先(下で上書き)
            ln_ref = style.find("a:lnRef", NS)
            fill_ref = style.find("a:fillRef", NS)
            if ln_ref is not None and ln_ref.get("idx") == "0":
                it["line"] = False
            elif ln_ref is not None:
                it["lncolor"] = self.pkg.color(ln_ref) or "black"
            if fill_ref is not None and fill_ref.get("idx") != "0":
                it["fill"] = self.pkg.color(fill_ref)
        if sppr is not None:
            g = sppr.find("a:prstGeom", NS)
            it["prst"] = g.get("prst") if g is not None else None
            if g is not None:
                for gd in g.iter("{%s}gd" % NS["a"]):
                    m = re.match(r"val (-?\d+)", gd.get("fmla") or "")
                    if m:
                        it["adj"][gd.get("name")] = int(m.group(1)) / 100000
            cg = sppr.find("a:custGeom", NS)
            if cg is not None:
                it["cust"] = cg
            if sppr.find("a:noFill", NS) is not None:
                it["fill"] = None
            elif sppr.find("a:solidFill", NS) is not None:
                it["fill"] = self.pkg.color(sppr.find("a:solidFill", NS)) or "#ffffff"
            ln = sppr.find("a:ln", NS)
            if ln is not None:
                it["line"] = True
                if ln.find("a:noFill", NS) is not None:
                    it["line"] = False
                it["lncolor"] = self.pkg.color(ln.find("a:solidFill", NS)) or it["lncolor"]
                pd = ln.find("a:prstDash", NS)
                it["dash"] = pd is not None and pd.get("val") not in (None, "solid")
                for k in ("headEnd", "tailEnd"):
                    e = ln.find(f"a:{k}", NS)
                    it[k[:4]] = e is not None and e.get("type") not in (None, "none")
        if kind == "cxnSp":
            for k in ("stCxn", "endCxn"):
                n = el.find(f"{self.q('nvCxnSpPr')}/{self.q('cNvCxnSpPr')}/a:{k}", NS)
                if n is not None:
                    it[k] = (n.get("id"), int(n.get("idx") or 0))
        if kind == "pic":
            blip = el.find(".//a:blip", NS)
            if blip is not None:
                it["image"] = self.pkg.media(self.base, self.rels, cast(str, blip.get("{%s}embed" % NS["r"])))
        if it["name"] in self.checks:
            it["check"] = self.checks[it["name"]]
        self.items.append(it)

    def emit_table(self, frame: ET.Element, box: Box) -> None:
        """graphicFrame の表(a:tbl)。セルを矩形 + 文字にする。結合セル(gridSpan / rowSpan)は左上のセルに広げる。"""
        tbl = frame.find(".//a:tbl", NS)
        if tbl is None:
            return
        widths = [int(cast(str, c.get("w"))) for c in tbl.findall("a:tblGrid/a:gridCol", NS)]
        rows_el = tbl.findall("a:tr", NS)
        heights = [int(cast(str, r.get("h"))) for r in rows_el]
        # 表の枠(graphicFrame の xfrm)に合わせて列幅・行高を比例させる
        sx = box[2] / sum(widths) if sum(widths) else 1 / EMU_PX
        sy = box[3] / sum(heights) if sum(heights) else 1 / EMU_PX
        xs = [box[0]]
        for w in widths:
            xs.append(xs[-1] + w * sx)
        ys = [box[1]]
        for h in heights:
            ys.append(ys[-1] + h * sy)
        fid = cast(ET.Element, frame.find(f"./*/{self.q('cNvPr')}", NS)).get("id")
        for r, tr in enumerate(rows_el):
            c = 0
            for tc in tr.findall("a:tc", NS):
                if c >= len(widths):
                    break  # tblGrid より多い tc(壊れた表)は無視
                if tc.get("hMerge") or tc.get("vMerge"):
                    c += 1
                    continue
                span, rspan = int(tc.get("gridSpan") or 1), int(tc.get("rowSpan") or 1)
                c1, r1 = min(c + span, len(xs) - 1), min(r + rspan, len(ys) - 1)
                pr = tc.find("a:tcPr", NS)
                fill = self.pkg.color(pr.find("a:solidFill", NS)) if pr is not None else None
                tf = text_fields(tc)
                tf["valign"] = {"ctr": "middle", "b": "bottom"}.get(cast(str, pr.get("anchor") if pr is not None else None), "top")  # 表は tcPr が縦位置を持つ
                self.items.append(dict(id=f"{fid}-{r}-{c}", name="", kind="sp", box=[xs[c], ys[r], xs[c1] - xs[c], ys[r1] - ys[r]],
                                       flipH=False, flipV=False, rot=0, prst="rect", cust=None, fill=fill, line=True, lncolor="black",
                                       dash=False, head=False, tail=False, adj={}, stCxn=None, endCxn=None, image=None, check=None, **tf))
                c += span

    def emit_chart(self, frame: ET.Element, box: Box) -> None:
        """graphicFrame のグラフ(c:chart)。枠と「グラフ: 題」だけの図形にし、値は chart に持たせる(chart_lines でテキストにする)。"""
        data_el = frame.find(".//a:graphicData", NS)
        ref = frame.find(".//c:chart", NS)
        if data_el is None or data_el.get("uri") != CHART_URI or ref is None:
            return
        xml = self.pkg.media(self.base, self.rels, cast(str, ref.get("{%s}id" % NS["r"])))
        data = chart_data(xml) if xml else None
        if data is None:
            return
        fid = cast(ET.Element, frame.find(f"./*/{self.q('cNvPr')}", NS)).get("id")
        label = chart_lines(data)[0]
        self.items.append(dict(id=fid, name="", kind="chart", box=box, flipH=False, flipV=False, rot=0, prst="rect", cust=None,
                               fill="#f4f4f4", line=True, lncolor="#808080", dash=False, head=False, tail=False, adj={},
                               stCxn=None, endCxn=None, image=None, check=None, chart=data, text=label, pt=12.0,
                               paras=[(label, 12.0, None)], align="center", valign="middle", wrap=True, vert=False, ins=(2.0, 1.0)))

    def emit_ole(self, frame: ET.Element, box: Box) -> None:
        """OLE オブジェクト(Excel のグラフなどを貼り付けたもの)。中身は描けないので、代わりに置かれている画像を枠いっぱいに貼る。

        その画像は mc:Fallback の中の p:pic にあり、walk は Fallback を飛ばすのでここで拾う。形式は EMF / WMF のことが多い。
        """
        data_el = frame.find(".//a:graphicData", NS)
        if data_el is None or not (data_el.get("uri") or "").endswith("/ole"):
            return
        blip = frame.find(".//a:blip", NS)
        image = self.pkg.media(self.base, self.rels, cast(str, blip.get("{%s}embed" % NS["r"]))) if blip is not None else None
        if not image:
            return
        fid = cast(ET.Element, frame.find(f"./*/{self.q('cNvPr')}", NS)).get("id")
        self.items.append(dict(id=fid, name="", kind="pic", box=box, flipH=False, flipV=False, rot=0, prst="rect", cust=None,
                               fill=None, line=False, lncolor="black", dash=False, head=False, tail=False, adj={},
                               stCxn=None, endCxn=None, image=image, check=None, **text_fields(frame)))

    def place(self, f: dict[str, Any] | None, box: Box | None, chbox: Box | None) -> list[float] | None:
        """図形の絶対枠 [x, y, w, h](px)。box: 親から与えられた枠(xlsx のアンカー)。chbox: グループの子座標系。"""
        if chbox is not None:
            box = cast(Box, box)
            f = cast(dict[str, Any], f)
            chx, chy, chcx, chcy = chbox
            sx = box[2] / chcx if chcx else 1
            sy = box[3] / chcy if chcy else 1
            return [box[0] + (f["x"] - chx) * sx, box[1] + (f["y"] - chy) * sy, f["cx"] * sx, f["cy"] * sy]
        if box is not None:
            b = list(box)
            # トップレベルのアンカー枠は回転後の外接枠。90°/270° なら幅と高さを入れ替えて回転前の枠にする
            if f and round(f["rot"]) % 180 == 90:
                cx, cy = b[0] + b[2] / 2, b[1] + b[3] / 2
                b = [cx - b[3] / 2, cy - b[2] / 2, b[3], b[2]]
            return b
        if f is None:
            return None
        return [f["x"] / EMU_PX, f["y"] / EMU_PX, f["cx"] / EMU_PX, f["cy"] / EMU_PX]  # pptx: xfrm が絶対座標

    def walk(self, el: ET.Element, box: Box | None, chbox: Box | None) -> None:
        """box: この要素群の絶対枠(pptx のトップレベルは None)。chbox: 子座標系の (x, y, cx, cy)。None なら子は自身の枠をそのまま使う。"""
        for ch in el:
            tag = ch.tag.split("}")[1]
            if tag in ("AlternateContent", "Choice", "Fallback"):
                if tag != "Fallback":
                    self.walk(ch, box, chbox)
                continue
            if tag == "graphicFrame":
                f = xfrm_node(ch.find(self.q("xfrm"), NS))  # graphicFrame は xfrm を直下に持つ(a: ではなく外側の名前空間)
                b = self.place(f, box, chbox)
                if b is not None:
                    self.emit_table(ch, b)
                    self.emit_chart(ch, b)
                    self.emit_ole(ch, b)
                continue
            if tag in ("sp", "cxnSp", "pic", "grpSp"):
                ph = ch.find(f"{self.q('nvSpPr')}/{self.q('nvPr')}/{self.q('ph')}", NS)
                if ph is not None and ph.get("type") in SKIP_PLACEHOLDERS:
                    continue
                pr = ch.find(self.q("grpSpPr") if tag == "grpSp" else self.q("spPr"), NS)
                f = xfrm(pr) if pr is not None else None
                if f is None and ph is not None and self.inherit:
                    f = self.inherit(ph)
                b = self.place(f, box, chbox)
                if b is None:
                    continue
                if tag == "grpSp":
                    self.walk(ch, b, f.get("ch") if f else None)
                else:
                    self.emit(ch, tag, b, f)


def collect(wb: Workbook, sheet: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]] | None:
    parts = wb.drawing_for(sheet)
    if parts is None:
        return None
    root, drawing, drels, ws_xml = parts
    cols, rows = cell_grid(wb.wb[sheet], wb.mdw)
    col = Collector("xdr", wb, drawing, drels, checks=wb.checkbox_states(ws_xml, wb.sheets[sheet]))

    def anchors(el: ET.Element) -> Iterator[ET.Element]:
        for a in el:
            tag = a.tag.split("}")[1]
            if tag.endswith("Anchor"):
                yield a
            elif tag == "AlternateContent":  # Choice の中にアンカーが入っていることがある(フォームコントロール)
                for choice in a:
                    if choice.tag.split("}")[1] == "Choice":
                        yield from anchors(choice)

    for a in anchors(root):
        col.walk(a, anchor_box(a, cols, rows), None)
    return col.items, cell_items(wb.wb[sheet], cols, rows, wb.theme)


def collect_slide(pres: Presentation, part: str) -> list[dict[str, Any]]:
    root = ET.fromstring(pres.z.read(part))
    rels = pres._rels(f"{Path(part).parent}/_rels/{Path(part).name}.rels")
    col = Collector("p", pres, part, rels, inherit=lambda ph: pres.placeholder_xfrm(part, ph))
    col.walk(cast(ET.Element, root.find("p:cSld/p:spTree", NS)), None, None)
    return col.items


# ---------------------------------------------------------------- connectors


def site_point(it: dict[str, Any], idx: int) -> tuple[Point, str]:
    """接続点番号 -> (点, 外向きの方向)。rect 系は 0=上 1=左 2=下 3=右、ellipse は 8 方位。"""
    x, y, w, h = it["box"]
    cx, cy = x + w / 2, y + h / 2
    if it["prst"] == "ellipse":
        ang = -math.pi / 2 + idx * math.pi / 4  # 0 が上、時計回り
        px, py = cx + w / 2 * math.cos(ang), cy + h / 2 * math.sin(ang)
        d = ("up", "up", "left", "down", "down", "down", "right", "up")[idx % 8]
        # 斜め方向はどちらかに寄せる
        if idx in (1, 7):
            d = "up"
        if idx in (3, 5):
            d = "down"
        return (px, py), d
    return [((cx, y), "up"), ((x, cy), "left"), ((cx, y + h), "down"), ((x + w, cy), "right")][idx % 4]


def free_ends(it: dict[str, Any]) -> tuple[Point, Point]:
    """flip と rot を反映した始点・終点。"""
    x, y, w, h = it["box"]
    x0, x1 = (x + w, x) if it["flipH"] else (x, x + w)
    y0, y1 = (y + h, y) if it["flipV"] else (y, y + h)
    p0, p1 = rotate_pts([(x0, y0), (x1, y1)], it["box"], it["rot"])
    return p0, p1


def connector_path(it: dict[str, Any], by_id: dict[str, dict[str, Any]]) -> list[Point]:
    """コネクタの折れ点列。"""
    p0, p1 = free_ends(it)
    d0 = d1 = None
    if it["stCxn"] and it["stCxn"][0] in by_id:
        p0, d0 = site_point(by_id[it["stCxn"][0]], it["stCxn"][1])
    if it["endCxn"] and it["endCxn"][0] in by_id:
        p1, d1 = site_point(by_id[it["endCxn"][0]], it["endCxn"][1])
    prst = it["prst"] or ""
    if not prst.startswith("bentConnector"):
        return [p0, p1]
    if d0 is None and d1 is None:
        return preset_bent(prst, it)
    return routed_bent(p0, d0, p1, d1)


def preset_bent(prst: str, it: dict[str, Any]) -> list[Point]:
    """接続情報なし: 枠内にプリセットどおり(2 は L、3 は Z、4 は 4 折れ)。flip と rot を反映する。"""
    x, y, w, h = it["box"]
    x0, x1 = (x + w, x) if it["flipH"] else (x, x + w)
    y0, y1 = (y + h, y) if it["flipV"] else (y, y + h)
    adj = it["adj"]
    if prst == "bentConnector2":
        pts = [(x0, y0), (x1, y0), (x1, y1)]
    elif prst == "bentConnector3":
        mx = x0 + (x1 - x0) * adj.get("adj1", 0.5)
        pts = [(x0, y0), (mx, y0), (mx, y1), (x1, y1)]
    else:
        ax = x0 + (x1 - x0) * adj.get("adj1", 0.4375)
        ay = y0 + (y1 - y0) * adj.get("adj2", 0.4375)
        pts = [(x0, y0), (ax, y0), (ax, ay), (x1, ay), (x1, y1)]
    return rotate_pts(pts, it["box"], it["rot"])


def _step(p: Point, d: str, dist: float) -> Point:
    x, y = p
    return {"up": (x, y - dist), "down": (x, y + dist), "left": (x - dist, y), "right": (x + dist, y)}[d]


def _vertical(d: str) -> bool:
    return d in ("up", "down")


def routed_bent(p0: Point, d0: str | None, p1: Point, d1: str | None) -> list[Point]:
    """接続点の向きから経路を組む。片側だけ接続されている場合は、その向きに出てから相手へ折れる。"""
    if d0 is None:
        pts = routed_bent(p1, d1, p0, None)
        return pts[::-1]
    if d1 is None:
        # 出る向きにまず進み、相手の座標へ 1 回折れる
        if _vertical(d0):
            return [p0, (p0[0], p1[1]), p1]
        return [p0, (p1[0], p0[1]), p1]
    opposite = {"up": "down", "down": "up", "left": "right", "right": "left"}
    if d1 == opposite[d0]:  # Z 型
        if _vertical(d0):
            my = (p0[1] + p1[1]) / 2
            return [p0, (p0[0], my), (p1[0], my), p1]
        mx = (p0[0] + p1[0]) / 2
        return [p0, (mx, p0[1]), (mx, p1[1]), p1]
    if d1 == d0:  # U 型: 両端から外へ出て並走
        if _vertical(d0):
            yy = (max if d0 == "down" else min)(p0[1], p1[1])
            yy = _step((0, yy), d0, ELBOW_MARGIN)[1]
            return [p0, (p0[0], yy), (p1[0], yy), p1]
        xx = (max if d0 == "right" else min)(p0[0], p1[0])
        xx = _step((xx, 0), d0, ELBOW_MARGIN)[0]
        return [p0, (xx, p0[1]), (xx, p1[1]), p1]
    # 直交: L 型
    if _vertical(d0):
        return [p0, (p0[0], p1[1]), p1]
    return [p0, (p1[0], p0[1]), p1]


# ---------------------------------------------------------------- drawing


class Canvas:
    def __init__(self, w: float, h: float, scale: float = SCALE) -> None:
        self.scale = scale  # 描画倍率。図形の座標(px, 96dpi)にこれを掛けて画素にする
        self.im = Image.new("RGB", (int(w * scale) + 1, int(h * scale) + 1), "white")
        self.d = ImageDraw.Draw(self.im)

    def s(self, v: float) -> float:
        return v * self.scale

    def pts(self, seq: Iterable[Point]) -> list[Point]:
        return [(self.s(x), self.s(y)) for x, y in seq]

    def polyline(self, seq: Iterable[Point], dash: bool=False, width: float=1, color: str="black") -> None:
        seq = self.pts(seq)
        if not dash:
            self.d.line(seq, fill=color, width=max(1, int(width * self.scale)))
            return
        for (x0, y0), (x1, y1) in zip(seq, seq[1:]):
            L = math.hypot(x1 - x0, y1 - y0)
            n = max(1, int(L / (6 * self.scale)))
            for i in range(0, n, 2):
                t0, t1 = i / n, min(1, (i + 1) / n)
                self.d.line([(x0 + (x1 - x0) * t0, y0 + (y1 - y0) * t0), (x0 + (x1 - x0) * t1, y0 + (y1 - y0) * t1)], fill=color, width=max(1, int(width * self.scale)))

    def arrowhead(self, tip: Point, frm: Point, size: float=7, color: str="black") -> None:
        tx, ty, fx, fy = self.s(tip[0]), self.s(tip[1]), self.s(frm[0]), self.s(frm[1])
        if (tx, ty) == (fx, fy):
            return
        ang = math.atan2(ty - fy, tx - fx)
        L = size * self.scale
        a = ang + math.radians(150)
        b = ang - math.radians(150)
        self.d.polygon([(tx, ty), (tx + L * math.cos(a), ty + L * math.sin(a)), (tx + L * math.cos(b), ty + L * math.sin(b))], fill=color)

    def text_in_box(self, text: str, box: Box, pt: float, align: str="center", color: str="black", valign: str="middle", wrap: bool=True,
                    paras: list[tuple[str, float, float | None]] | None=None, vert: bool=False, ins: Point=(2.0, 1.0)) -> None:
        """paras: 段落ごとの (文字, pt, 行送り pt or None)。None なら text を pt で。vert: 縦書き(1 文字 1 行で積む)。"""
        if not text:
            return
        if paras is None:
            paras = [(p, pt, None) for p in text.split("\n")]
        paras = [(p.translate(GLYPH_FALLBACK), ppt, plh) for p, ppt, plh in paras]
        if vert:
            paras = [(ch, ppt, None) for p, ppt, _ in paras for ch in p]
            wrap, align = False, "center"
        x, y, w, h = [self.s(v) for v in box]
        mx, my = self.s(ins[0]), self.s(ins[1])
        maxw = max(w - 2 * mx, 10) if wrap else 10**9
        lines = []  # (行, フォント, 行送り px)
        for para, ppt, plh in paras:
            f = font(ppt, self.scale)
            lh = self.s(plh * 96 / 72) if plh else f.size * 1.15
            cur = ""
            for ch in para:
                if self.d.textlength(cur + ch, font=f) > maxw and cur:
                    lines.append((cur, f, lh))
                    cur = ch
                else:
                    cur += ch
            lines.append((cur, f, lh))
        total = sum(lh for _, _, lh in lines)
        if valign == "middle":
            ty = y + max(0, (h - total) / 2)
        elif valign == "bottom":
            ty = y + max(0, h - total - my)
        else:
            ty = y + my
        for line, f, lh in lines:
            tw = self.d.textlength(line, font=f)
            tx = x + (w - tw) / 2 if align == "center" else (x + mx if align == "left" else x + w - tw - mx)
            self.d.text((tx, ty), line, font=f, fill=color)
            ty += lh


def rotate_pts(seq: list[Point], box: Box, rot: float) -> list[Point]:
    if not rot:
        return seq
    x, y, w, h = box
    cx, cy = x + w / 2, y + h / 2
    a = math.radians(rot)
    out = []
    for px, py in seq:
        dx, dy = px - cx, py - cy
        out.append((cx + dx * math.cos(a) - dy * math.sin(a), cy + dx * math.sin(a) + dy * math.cos(a)))
    return out


def shape_outline(it: dict[str, Any]) -> list[Point] | None:
    """プリセット図形の輪郭(多角形)。曲線は折れ線で近似。None なら矩形。"""
    x, y, w, h = it["box"]
    p = it["prst"]
    if it["cust"] is not None:
        return cust_outline(it)
    if p in ("ellipse", "flowChartConnector"):
        return [(x + w / 2 + w / 2 * math.cos(t), y + h / 2 + h / 2 * math.sin(t)) for t in [i * math.pi / 18 for i in range(36)]]
    if p == "flowChartDecision":
        return [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
    if p == "flowChartOffpageConnector":
        return [(x, y), (x + w, y), (x + w, y + h * 0.8), (x + w / 2, y + h), (x, y + h * 0.8)]
    if p == "flowChartDocument":
        pts = [(x, y), (x + w, y), (x + w, y + h * 0.83)]
        for i in range(1, 12):
            t = i / 12
            pts.append((x + w * (1 - t), y + h * (0.83 + 0.17 * math.sin(t * 2 * math.pi) * 0.5)))
        pts.append((x, y + h * 0.83))
        return pts
    if p == "rightArrow":
        return [(x, y + h * 0.25), (x + w * 0.65, y + h * 0.25), (x + w * 0.65, y), (x + w, y + h / 2), (x + w * 0.65, y + h), (x + w * 0.65, y + h * 0.75), (x, y + h * 0.75)]
    if p == "upArrow":
        return [(x + w * 0.25, y + h), (x + w * 0.25, y + h * 0.35), (x, y + h * 0.35), (x + w / 2, y), (x + w, y + h * 0.35), (x + w * 0.75, y + h * 0.35), (x + w * 0.75, y + h)]
    if p == "bentUpArrow":
        return [(x, y + h * 0.7), (x + w * 0.6, y + h * 0.7), (x + w * 0.6, y + h * 0.3), (x + w * 0.45, y + h * 0.3), (x + w * 0.75, y), (x + w + 0, y + h * 0.3), (x + w * 0.9, y + h * 0.3), (x + w * 0.9, y + h), (x, y + h)]
    if p == "roundRect":
        r = min(w, h) * 0.15
        pts = []
        for (cx, cy, a0) in [(x + w - r, y + r, -90), (x + w - r, y + h - r, 0), (x + r, y + h - r, 90), (x + r, y + r, 180)]:
            for i in range(7):
                t = math.radians(a0 + i * 15)
                pts.append((cx + r * math.cos(t), cy + r * math.sin(t)))
        return pts
    return None


def cust_outline(it: dict[str, Any]) -> list[Point] | None:
    x, y, w, h = it["box"]
    path = it["cust"].find("a:pathLst/a:path", NS)
    if path is None:
        return None
    pw, ph = int(path.get("w") or 1) or 1, int(path.get("h") or 1) or 1
    pts = []
    for node in path:
        for pt in node.iter("{%s}pt" % NS["a"]):
            pts.append((x + int(cast(str, pt.get("x"))) / pw * w, y + int(cast(str, pt.get("y"))) / ph * h))
    return pts or None


def draw_shape(c: Canvas, it: dict[str, Any]) -> None:
    x, y, w, h = it["box"]
    p = it["prst"] or ""
    if it["check"] is not None:
        draw_checkbox(c, it)
        return
    if it["image"]:
        try:
            im = open_image(it["image"]).convert("RGBA")
            im = im.resize((max(1, int(c.s(w))), max(1, int(c.s(h)))))
            c.im.paste(im, (int(c.s(x)), int(c.s(y))), im)
        except Exception:
            c.d.rectangle(c.pts([(x, y), (x + w, y + h)]), outline="gray")
        return
    if p == "line":
        if not it["line"]:
            return
        p0, p1 = free_ends(it)
        c.polyline([p0, p1], dash=it["dash"], color=it["lncolor"], width=1.5)
        if it["tail"]:
            c.arrowhead(p1, p0, size=9, color=it["lncolor"])
        if it["head"]:
            c.arrowhead(p0, p1, size=9, color=it["lncolor"])
        return
    if p in ("can", "flowChartMagneticDisk", "flowChartOnlineStorage"):
        draw_cylinder(c, it)
    else:
        outline = shape_outline(it)
        if outline is None:
            outline = [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]
        outline = rotate_pts(outline, it["box"], it["rot"])
        closed = it["cust"] is None or it["cust"].find(".//a:close", NS) is not None
        if closed:
            c.d.polygon(c.pts(outline), fill=it["fill"], outline=it["lncolor"] if it["line"] else None)
        else:
            c.polyline(outline, dash=it["dash"], color=it["lncolor"])
    if p in ("borderCallout1", "wedgeRectCallout"):
        # 引き出し線: 枠の左下から少し外へ
        c.polyline([(x, y + h), (x - w * 0.15, y + h * 1.5)])
    c.text_in_box(it["text"], it["box"], it["pt"], align=it["align"], color=it.get("color", "black"), valign=it["valign"],
                  wrap=it["wrap"], paras=it["paras"], vert=it["vert"], ins=it["ins"])


def draw_cylinder(c: Canvas, it: dict[str, Any]) -> None:
    x, y, w, h = it["box"]
    ry = min(h * 0.18, w * 0.25)
    body = [(x, y + ry), (x, y + h - ry)]
    c.d.rectangle(c.pts([(x, y + ry), (x + w, y + h - ry)]), fill=it["fill"], outline=None)
    c.polyline([(x, y + ry), (x, y + h - ry)])
    c.polyline([(x + w, y + ry), (x + w, y + h - ry)])
    c.d.ellipse(c.pts([(x, y + h - 2 * ry), (x + w, y + h)]), fill=it["fill"], outline="black")
    c.d.rectangle(c.pts([(x + 1, y + h - 2 * ry), (x + w - 1, y + h - ry)]), fill=it["fill"])
    c.d.ellipse(c.pts([(x, y), (x + w, y + 2 * ry)]), fill=it["fill"], outline="black")


def draw_checkbox(c: Canvas, it: dict[str, Any]) -> None:
    x, y, w, h = it["box"]
    size = min(h * 0.7, 10)
    bx, by = x + 1, y + (h - size) / 2
    state = it["check"]
    fill = "#c8c8c8" if state == "Mixed" else "white"
    c.d.rectangle(c.pts([(bx, by), (bx + size, by + size)]), fill=fill, outline="black")
    if state == "Checked":
        c.polyline([(bx, by), (bx + size, by + size)])
        c.polyline([(bx + size, by), (bx, by + size)])
    c.text_in_box(it["text"], [bx + size + 3, y, w - size - 4, h], it["pt"], align="left")


def draw_connector(c: Canvas, it: dict[str, Any], by_id: dict[str, dict[str, Any]]) -> None:
    if not it["line"]:
        return
    pts = connector_path(it, by_id)
    c.polyline(pts, dash=it["dash"], color=it["lncolor"], width=1.5)
    if it["tail"]:
        c.arrowhead(pts[-1], pts[-2], size=9, color=it["lncolor"])
    if it["head"]:
        c.arrowhead(pts[0], pts[1], size=9, color=it["lncolor"])


def render(items: list[dict[str, Any]], cells: list[dict[str, Any]], title: str, size: tuple[float, float] | None = None) -> Image.Image:
    """size: 画布の大きさ(px)。None なら図形の範囲に合わせる(xlsx)。pptx はスライドの大きさを渡す。"""
    top = TITLE_PX
    if size:
        W, H = size[0], size[1] + top
    else:
        xs = [i["box"][0] + i["box"][2] for i in items]
        ys = [i["box"][1] + i["box"][3] for i in items]
        W, H = max(xs) + 20, max(ys) + 20 + top
    c = Canvas(W, H)
    c.d.text((c.s(8), c.s(6)), title, font=font(10, c.scale), fill="black")
    by_id = {i["id"]: i for i in items}
    for it in items:
        it["box"] = [it["box"][0], it["box"][1] + top, it["box"][2], it["box"][3]]
    # セル(下地): 図形の範囲内だけ描く
    for ce in cells:
        x, y, w, h = ce["box"]
        y += top
        if x > W or y > H:
            continue
        if ce["fill"]:
            c.d.rectangle(c.pts([(x, y), (x + w, y + h)]), fill=ce["fill"])
        c.text_in_box(ce["text"], [x, y, w, h], ce["pt"], align=ce["align"], valign="top", wrap=ce["wrap"])
    # XML の出現順 = Excel の重なり順。後の図形が上に来る(白い箱で下の図形を隠す表現に依存した図がある)
    for it in items:
        if it["kind"] == "cxnSp":
            draw_connector(c, it, by_id)
        else:
            draw_shape(c, it)
    return c.im


def render_region(items: list[dict[str, Any]], region: tuple[float, float, float, float], scale: float) -> Image.Image:
    """図形のうち region(x, y, w, h。px, 96dpi)に掛かるものだけを、region の範囲で scale 倍に描く(題もセルも描かない)。

    xlsx に貼られた画像を、上に重ねた図形(番号・枠・吹き出し・矢印)ごと、画像の元の解像度に近い倍率で読ませるため。
    items の box は書き換えない(複製してずらす)。
    """
    rx, ry, rw, rh = region
    shifted = [{**it, "box": [it["box"][0] - rx, it["box"][1] - ry, it["box"][2], it["box"][3]]} for it in items if it["box"]]
    by_id = {i["id"]: i for i in shifted}
    c = Canvas(rw, rh, scale)
    for it in shifted:
        x, y, w, h = it["box"]
        if x > rw or y > rh or x + w < 0 or y + h < 0:
            continue
        if it["kind"] == "cxnSp":
            draw_connector(c, it, by_id)
        else:
            draw_shape(c, it)
    return c.im


def render_sheets(path: Path, only: str | None = None) -> dict[str, Image.Image]:
    """図形が MIN_SHAPES 以上あるシートを描いて、シート名 → 画像 で返す。"""
    wb = Workbook(path)
    images = {}
    for sheet in wb.sheets:
        if only and sheet != only:
            continue
        got = collect(wb, sheet)
        if not got or len(got[0]) < MIN_SHAPES:
            continue
        items, cells = got
        images[sheet] = render(items, cells, f"{path.stem} / {sheet}")
    return images


def render_slides(path: Path, only: str | None = None) -> dict[str, Image.Image]:
    """図形が MIN_SHAPES 以上あるスライドを描いて、スライド名("slide01")→ 画像 で返す。"""
    pres = Presentation(path)
    images = {}
    for name, part in pres.slides:
        if only and name != only:
            continue
        items = collect_slide(pres, part)
        if len(items) < MIN_SHAPES:
            continue
        images[name] = render(items, [], f"{path.stem} / {name}", size=pres.size)
    return images


def render_file(path: Path, only: str | None = None) -> dict[str, Image.Image]:
    return (render_slides if path.suffix == ".pptx" else render_sheets)(path, only)


def render_workbook(path: Path, out_dir: Path, only: str | None = None) -> list[Path]:
    written = []
    for sheet, im in render_file(path, only).items():
        out = out_dir / f"{path.stem}__{re.sub(r'[\\/:*?\"<>|]', '_', sheet)}.png"
        out_dir.mkdir(parents=True, exist_ok=True)
        im.save(out)
        written.append(out)
    return written


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", help="xlsx / pptx")
    ap.add_argument("--sheet", help="シート名、または pptx なら slide01 のようなスライド名")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    if not FONT_PATH:
        print("日本語フォントが見つかりません(fonts-ipafont-gothic を入れてください)", file=sys.stderr)
    for f in args.files:
        for p in render_workbook(Path(f), Path(args.out), args.sheet):
            print(p)
