"""段階 3 の抽出器(structured)の xlsx 部分。openpyxl でセルを読み、行ごとに 1 行のテキストにする。

帳票レイアウトの Excel 向けの自前実装。シート全体を 1 つの表としては読まず、行の並びを残す。そのうえで:

- 表(ヘッダ行 + データ行)を見つけたら、データ行を `列名: 値 | 列名: 値` の形にする。
  列はヘッダの結合セルの範囲で対応させるので、空欄があっても値が別の列にずれない
  (テーブル定義書の PK と必須、テスト仕様書の大項目と中項目など)。
  左側の階層の列(1 列目の次から最初に値のある列まで)が空欄なら、上の行の値を引き継ぐ。
  文書の種類ごとの作り込みはせず、どのテンプレートにも同じ規則を当てる
- 番号の並びだけの行(表のヘッダ下の `1 2 3 4 5 6 7`)、値の無いラベル(`作成者：`)、括弧だけの枠
  (チェックボックスの残り)を落とす
- テンプレートの名前は決め打ちせず、繰り返しと参照で見分けて落とす:
  - ヘッダ枠: シートの先頭の行が、ブックの半分を超えるシートの先頭にも同じ文字である(`PJ名 | … | 成果物名 | …`)
  - 目次: 題の行を除いた行の半分を超えて、ほかのシートの見出しかシート名と同じ(空白の違いは無視する)
  - 入力規則の選択肢の置き場: 入力規則(プルダウン)から参照されているシート

画面遷移図などの Excel 図形内のテキストは openpyxl では読めないので、drawing XML から別途拾う。
図形の層に置かれたグラフは、グラフの XML から題と系列ごとの値を行にする(render_drawing.chart_lines)。
"""

import re
import warnings
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import cast
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.workbook.workbook import Workbook as ExcelWorkbook
from openpyxl.worksheet.worksheet import Worksheet

from . import Chunked, Line, Section
from .chunking import chunk_sections
from .render_drawing import NS as DRAWING_NS
from .render_drawing import Package, chart_data, chart_lines

FRAME_ROWS = 5  # ヘッダ枠を探す、シートの先頭からの行数(値のある行で数える)
HEADING_RE = re.compile(r"^(\(\d+(-\d+)?\)|(\d+\.)+)\s*\S")  # "2.1. 画面レイアウト", "(1) バリデーション処理"

warnings.filterwarnings("ignore", module="openpyxl")  # Data Validation extension の警告


def cell_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


NUMBER_RE = re.compile(r"[\d.\-]+")
SHEET_REF_RE = re.compile(r"^=?'?([^'!]+)'?!")  # 入力規則の参照先 `データ!$A$1:$A$5`、`'シート 1'!$A:$A`
BRACKETS_RE = re.compile(r"[()（）\[\]［］【】\s]*")
MAX_LABEL_CHARS = 20  # ヘッダのラベルとみなす長さ


Cell = tuple[int, int, str, int]  # (開始列, 終了列, テキスト, 縦に結合している行数)


def row_cells(ws: Worksheet) -> list[tuple[int, list[Cell]]]:
    """行ごとに、値のあるセルを返す。結合セルは右端の列と縦の行数も持つ。"""
    spans = {(m.min_row, m.min_col): (m.max_col, m.max_row - m.min_row + 1) for m in ws.merged_cells.ranges}
    rows = []
    for row in ws.iter_rows():
        cells = []
        for cell in row:
            if cell.value in (None, ""):
                continue
            text = cell_text(cell.value)
            if text:
                end, height = spans.get((cell.row, cell.column), (cell.column, 1))
                cells.append((cell.column, end, text, height))
        if cells:
            rows.append((row[0].row, cells))
    return rows


def clean(cells: list[Cell]) -> list[Cell]:
    """ノイズのセルを落とす: 値の無いラベル(`作成者：`)、括弧だけの枠、補助列で横に繰り返された同じ値。"""
    out: list[Cell] = []
    for c in cells:
        text = c[2]
        if text.endswith(("：", ":")) or BRACKETS_RE.fullmatch(text):
            continue
        if out and out[-1][2] == text:
            continue
        out.append(c)
    return out


def is_noise_row(cells: list[Cell]) -> bool:
    """番号の並びだけの行(表のヘッダ下の `1 2 3 4 5 6 7`)。数値だけの表の行(`2022 | 120 | 100`)は残す。"""
    texts = [c[2] for c in cells]
    if len(texts) < 3 or not all(t.isdigit() for t in texts):
        return False
    return all(int(b) == int(a) + 1 for a, b in zip(texts, texts[1:]))


Row = tuple[int, list[Cell]]


def row_text(cells: list[Cell]) -> str:
    return " | ".join(c[2] for c in cells)


def strip_frames(sheets: dict[str, list[Row]]) -> dict[str, list[Row]]:
    """各シートの先頭から、ヘッダ枠の行を落とす。

    ヘッダ枠は、ブックの半分を超える(2 シート以上の)シートで、先頭 FRAME_ROWS 行の中に同じ文字で出る行。
    半分ちょうどは含めない(中身のシートが 1 枚のブックでは、その見出しが目次にも出る)。
    先頭から続く間だけ落とすので、枠の後ろにある表のヘッダがたまたま同じでも残る。
    """
    counts = Counter(text for rows in sheets.values() for text in {row_text(cells) for _, cells in rows[:FRAME_ROWS]})
    frame = {t for t, n in counts.items() if n >= 2 and n > len(sheets) / 2}
    out = {}
    for name, rows in sheets.items():
        i = 0
        while i < len(rows) and row_text(rows[i][1]) in frame:
            i += 1
        out[name] = rows[i:]
    return out


def squash(text: str) -> str:
    """空白とセルの区切りを除く(目次とシートの見出しで、空白やセルの分け方が違うことがある)。"""
    return re.sub(r"[\s|]+", "", text)


def toc_sheets(sheets: dict[str, list[Row]]) -> set[str]:
    """目次のシート: 題の行を除いた行の半分を超えて、ほかのシートの見出しかシート名と同じ。

    行ごとに、行全体か、セルがどれも見出しと同じなら一致とする(2 段組みの目次は 1 行に見出しが 2 つ並ぶ)。
    実物の目次は更新漏れや誤字でずれることがあるので、全行の一致は求めない。
    見出しとだけ比べるので、「なし」の 1 行や空の表のヘッダがほかのシートにあるだけのシートは残る。
    どのシートにも出る小見出し(`(1) 概要`)は目次の項目にならないので、(そのシート以外の)2 シート以上に出る見出しとは
    比べない。
    """
    headings = {
        name: {squash(name)}
        | {squash(row_text(cells)) for _, cells in rows if HEADING_RE.match(" ".join(c[2] for c in cells))}
        for name, rows in sheets.items()
    }
    counts = Counter(h for hs in headings.values() for h in hs)
    found = set()
    for name, rows in sheets.items():
        own = headings[name]
        others = {h for other, hs in headings.items() if other != name for h in hs if counts[h] - (h in own) == 1}
        body = rows[1:]
        hits = sum(
            squash(row_text(cells)) in others or all(squash(c[2]) in others for c in cells) for _, cells in body
        )
        if body and hits > len(body) / 2:
            found.add(name)
    return found


def validation_sheets(wb: ExcelWorkbook) -> set[str]:
    """入力規則(プルダウン)の選択肢の置き場になっているシート。参照は直接(`データ!$A$1:$A$5`)か名前の定義経由。"""
    names = {n: d.attr_text for n, d in wb.defined_names.items()}
    for ws in wb.worksheets:
        names |= {n: d.attr_text for n, d in ws.defined_names.items()}
    found = set()
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            ref = (dv.formula1 or "").lstrip("=")
            m = SHEET_REF_RE.match(names.get(ref, ref) or "")
            if m and m.group(1) in wb.sheetnames and m.group(1) != ws.title:
                found.add(m.group(1))
    return found


def is_label_row(cells: list[Cell]) -> bool:
    """ラベルだけの行か: どのセルも短く、数字だけのものが無い。"""
    return all(len(c[2]) <= MAX_LABEL_CHARS and not NUMBER_RE.fullmatch(c[2]) for c in cells)


Header = list[tuple[int, int, str]]  # (開始列, 終了列, ラベル)


def header_block(rows: list[Row], i: int) -> tuple[Header, int] | None:
    """rows[i] から始まるヘッダを読む。(ヘッダ, ヘッダの次の行の添字) を返す。ヘッダでなければ None。

    ヘッダは短いラベルが 3 つ以上並ぶ行。セルが縦に結合していれば、その高さの分の行もヘッダの段として読み
    (画面設計書の「表示情報」の下に「画面項目名 | 画面項目種別 | …」がある形)、列ごとに一番下の段のラベルを使う。
    """
    row_no, cells = rows[i]
    if len(cells) < 3 or not is_label_row(cells):
        return None
    height = max(c[3] for c in cells)
    block = [cells]
    j = i + 1
    while j < len(rows) and rows[j][0] < row_no + height and is_label_row(rows[j][1]):
        block.append(rows[j][1])
        j += 1
    # 下の段にラベルがある列は、下の段のラベルを使う(上の段は下の段を束ねる見出し)
    labels: Header = []
    for depth, level in enumerate(block):
        for start, end, text, _ in level:
            covered = any(s <= end and start <= e for lower in block[depth + 1 :] for s, e, _, _ in lower)
            if not covered:
                labels.append((start, end, text))
    labels.sort()
    return labels, j


def assign(header: Header, cells: list[Cell], nearest: bool = False) -> tuple[dict[str, str], list[str]]:
    """セルをヘッダのラベルに割り当てる。(ラベル → 値, どのラベルにも入らなかった値)。

    nearest なら、ラベルの範囲外のセルは左側で一番近いラベルに付ける(ヘッダが結合されておらず、
    値だけ右の列にはみ出している表がある)。左にラベルが無ければ、どのラベルにも入らない値として返す。
    """
    values: dict[str, str] = {}
    missed: list[str] = []
    for col, _, text, _ in cells:
        label = next((lab for start, end, lab in header if start <= col <= end), None)
        if label is None and nearest:
            label = next((lab for start, _, lab in reversed(header) if start <= col), None)
        if label is None:
            missed.append(text)
        else:
            values[label] = f"{values[label]} {text}" if label in values else text
    return values, missed


def is_data_row(header: Header, cells: list[Cell]) -> bool:
    values, missed = assign(header, cells)
    return len(values) >= 1 and len(missed) <= len(cells) // 5


def confirms_header(header: Header, following: list[list[Cell]]) -> bool:
    """ヘッダ候補の直後 2 行が表のデータ行か。

    どちらも 2 つ以上の列に値が入り、どこかにラベルらしくない値(数字、長い文)があること。
    帳票の「項目名 | 値 | 項目名」の行(`授受方式 | HTTP | ﾌｨｰﾙﾄﾞｾﾊﾟﾚｰﾀ`)を表と取り違えないため。
    """
    if len(following) < 2:
        return False
    if not all(len(assign(header, cells)[0]) >= 2 and is_data_row(header, cells) for cells in following):
        return False
    return not all(is_label_row(cells) for cells in following)


def sheet_rows(ws: Worksheet) -> list[Row]:
    """値のある行を、ノイズのセルと番号の並びの行を落として返す。"""
    rows = [(no, clean(cells)) for no, cells in row_cells(ws)]
    return [(no, cells) for no, cells in rows if cells and not is_noise_row(cells)]


def sheet_lines(rows: list[Row]) -> list[Line]:
    """行ごとに (テキスト, 種別) を返す。表のデータ行は `列名: 値 | …` にする。"""
    lines: list[Line] = []
    header: Header | None = None
    previous: dict[str, str] = {}
    i = 0
    while i < len(rows):
        cells = rows[i][1]
        texts = [c[2] for c in cells]
        if len(texts) == 1 and HEADING_RE.match(texts[0]):
            header = None
            lines.append((texts[0], "heading"))
            i += 1
            continue
        if header and is_data_row(header, cells):
            values, missed = assign(header, cells, nearest=True)
            labels = [lab for _, _, lab in header]
            filled = [k for k, lab in enumerate(labels) if lab in values]
            first = next((k for k in filled if k > 0), 1)
            for k in range(1, first):  # 左側の階層の列の空欄は上の行を引き継ぐ
                if labels[k] in previous:
                    values[labels[k]] = previous[labels[k]]
            previous = values
            # ヘッダにラベルの無い列の値も落とさず、ラベル無しで後ろに付ける
            lines.append((" | ".join([f"{lab}: {values[lab]}" for lab in labels if lab in values] + missed), ""))
            i += 1
            continue
        found = header_block(rows, i)
        if found and confirms_header(found[0], [c for _, c in rows[found[1] : found[1] + 2]]):
            header, previous = found[0], {}
            i = found[1]  # ヘッダ行はデータ行のラベルになるので、単独の行としては出さない
            continue
        header = None
        lines.append((" | ".join(texts), ""))
        i += 1
    return lines


def shape_texts(path: Path) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """(シート名 → 図形内テキストのリスト, シート名 → グラフの行のリスト)。

    関係ファイル(rels)の読み方とパスの解決は render_drawing.Package に任せる(Excel 以外のツールが書いた xlsx は、
    属性の順番や Target の書き方(絶対パス)が違う)。
    """
    pkg = Package(path)
    result, charts = {}, {}
    sheet_rels = pkg._rels("xl/_rels/workbook.xml.rels")
    for m in re.finditer(r"<sheet [^>]*/>", pkg.z.read("xl/workbook.xml").decode()):
        name = re.search(r' name="([^"]+)"', m.group(0))
        rid = re.search(r' r:id="([^"]+)"', m.group(0))
        if not (name and rid and rid.group(1) in sheet_rels):
            continue
        sheet_part = Package._resolve("xl/workbook.xml", sheet_rels[rid.group(1)])
        texts, chart_rows = [], []
        for target in pkg._rels(f"{Path(sheet_part).parent}/_rels/{Path(sheet_part).name}.rels").values():
            if "drawings/drawing" not in target:
                continue
            drawing = Package._resolve(sheet_part, target)
            if drawing not in pkg.names:
                continue
            root = ET.fromstring(pkg.z.read(drawing))
            texts += [t.text.strip() for t in root.iter("{%s}t" % DRAWING_NS["a"]) if t.text and t.text.strip()]
            drawing_rels = pkg._rels(f"{Path(drawing).parent}/_rels/{Path(drawing).name}.rels")
            for ref in root.iter("{%s}chart" % DRAWING_NS["c"]):
                chart_part = Package._resolve(drawing, drawing_rels.get(cast(str, ref.get("{%s}id" % DRAWING_NS["r"])), ""))
                data = chart_data(pkg.z.read(chart_part)) if chart_part in pkg.names else None
                if data:
                    chart_rows += chart_lines(data)
        sheet = name.group(1).replace("&amp;", "&")
        if texts:
            result[sheet] = texts
        if chart_rows:
            charts[sheet] = chart_rows
    return result, charts


def sections(path: Path) -> list[Section]:
    """切る前のシートごとの行(vision もこれに図の説明を足して使う)。"""
    shapes, charts = shape_texts(path)
    wb = openpyxl.load_workbook(path, data_only=True)  # 結合セルの範囲を読むので read_only にしない
    sheets = strip_frames({ws.title: sheet_rows(ws) for ws in wb.worksheets})
    skip = toc_sheets(sheets) | validation_sheets(wb)
    sections = []
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        lines = sheet_lines(sheets[ws.title])
        if ws.title in shapes:
            lines.append(("図形テキスト: " + " / ".join(shapes[ws.title]), ""))
        lines += [(row, "") for row in charts.get(ws.title, [])]
        sections.append((ws.title, lines))
    return sections


def extract(path: Path) -> list[Chunked]:
    return chunk_sections(sections(path))
